from __future__ import annotations

import csv
import json
import math
import shutil
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

SUPPORTED_MODEL_SUFFIXES = {".gguf", ".npz", ".nrlmodel"}
SUPPORTED_REPORT_SUFFIXES = {".json", ".csv", ".png", ".pdf", ".docx", ".pptx"}
OBSERVATION_SIZE = 36
ACTION_SIZE = 2
ARCHITECTURES = {"residual_rnn", "continuous_time", "equilibrium"}


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, TypeError):
            pass
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def inspect_npz(path: Path) -> dict[str, Any]:
    import numpy as np

    with np.load(path, allow_pickle=False) as archive:
        if "parameters" not in archive.files:
            raise ValueError("NPZ must contain a 'parameters' array.")
        parameters = np.asarray(archive["parameters"])
    if parameters.ndim != 1:
        raise ValueError(f"Expected a flat parameter vector, got shape {parameters.shape}.")
    if not np.issubdtype(parameters.dtype, np.number):
        raise ValueError("Parameter vector must be numeric.")
    if not np.isfinite(parameters).all():
        raise ValueError("Parameter vector contains NaN or infinity.")
    return {
        "parameter_count": int(parameters.size),
        "parameter_dtype": str(parameters.dtype),
        "parameter_shape": list(parameters.shape),
    }


def export_nrlmodel(
    checkpoint: Path,
    destination: Path,
    metadata_path: Path | None = None,
) -> Path:
    """Wrap a native NeuroRaceLab flat-parameter NPZ in a validated .nrlmodel ZIP."""
    import numpy as np

    checkpoint = Path(checkpoint)
    destination = Path(destination)
    info = inspect_npz(checkpoint)
    metadata: dict[str, Any] = {}
    if metadata_path and Path(metadata_path).is_file():
        raw = load_json(Path(metadata_path))
        if isinstance(raw, dict):
            metadata = _json_safe(raw)
    architecture = metadata.get("architecture", checkpoint.name.split("_seed")[0])
    if architecture not in ARCHITECTURES:
        architecture = "unknown"
    hidden_size = metadata.get("hidden_size")
    manifest = {
        "format": "neuroracelab-controller",
        "format_version": 1,
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "architecture": architecture,
        "hidden_size": hidden_size,
        "observation": {
            "size": OBSERVATION_SIZE,
            "description": "32 ray distances, speed, slip angle, yaw rate, current steering",
        },
        "action": {
            "size": ACTION_SIZE,
            "description": "steering and longitudinal throttle/brake, each bounded to [-1, 1]",
        },
        **info,
        "source_checkpoint": checkpoint.name,
        "metadata_file_included": bool(metadata),
    }
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="nrlmodel-") as temp:
        tempdir = Path(temp)
        copied_npz = tempdir / "parameters.npz"
        with np.load(checkpoint, allow_pickle=False) as archive:
            np.savez_compressed(copied_npz, parameters=np.asarray(archive["parameters"]))
        (tempdir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        (tempdir / "metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
        with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
            bundle.write(tempdir / "manifest.json", "manifest.json")
            bundle.write(copied_npz, "parameters.npz")
            bundle.write(tempdir / "metadata.json", "metadata.json")
    validate_nrlmodel(destination)
    return destination


def validate_nrlmodel(path: Path) -> dict[str, Any]:

    path = Path(path)
    if not zipfile.is_zipfile(path):
        raise ValueError("Not a valid .nrlmodel ZIP bundle.")
    with zipfile.ZipFile(path, "r") as bundle:
        names = set(bundle.namelist())
        if not {"manifest.json", "parameters.npz"}.issubset(names):
            raise ValueError("Bundle must contain manifest.json and parameters.npz.")
        manifest = json.loads(bundle.read("manifest.json").decode("utf-8"))
        if manifest.get("format") != "neuroracelab-controller":
            raise ValueError("Unrecognised model bundle format.")
        if manifest.get("format_version") != 1:
            raise ValueError(f"Unsupported .nrlmodel version: {manifest.get('format_version')!r}")
        with tempfile.TemporaryDirectory(prefix="nrl-validate-") as temp:
            extracted = Path(temp) / "parameters.npz"
            extracted.write_bytes(bundle.read("parameters.npz"))
            info = inspect_npz(extracted)
    if info["parameter_count"] != manifest.get("parameter_count"):
        raise ValueError("Parameter count does not match the bundle manifest.")
    if manifest.get("observation", {}).get("size") != OBSERVATION_SIZE:
        raise ValueError("Unsupported observation schema.")
    if manifest.get("action", {}).get("size") != ACTION_SIZE:
        raise ValueError("Unsupported action schema.")
    return manifest


def import_model_file(source: Path, library_dir: Path) -> Path:
    """Copy supported model file into the library, validating native formats."""
    source, library_dir = Path(source), Path(library_dir)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_MODEL_SUFFIXES:
        raise ValueError(f"Unsupported model format: {suffix or '(no extension)'}")
    if not source.is_file():
        raise FileNotFoundError(source)
    if suffix == ".npz":
        inspect_npz(source)
    elif suffix == ".nrlmodel":
        validate_nrlmodel(source)
    elif suffix == ".gguf":
        # GGUF structural/runtime validation belongs to a GGUF-capable runtime.
        if source.stat().st_size < 8:
            raise ValueError("File is too small to be a valid GGUF model.")
        with source.open("rb") as handle:
            if handle.read(4) != b"GGUF":
                raise ValueError("File does not start with the GGUF magic header.")
    library_dir.mkdir(parents=True, exist_ok=True)
    destination = library_dir / source.name
    if source.resolve() != destination.resolve():
        shutil.copy2(source, destination)
    return destination


def load_report_data(source: Path) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    source = Path(source)
    if source.suffix.lower() == ".json":
        data = load_json(source)
        if isinstance(data, dict):
            records = (data.get("episodes") or data.get("history") or data.get("results") or data.get("generations") or data.get("metrics") or data.get("records"))
            if isinstance(records, list) and all(isinstance(x, dict) for x in records):
                return source.stem, records, data
            # Flatten top-level scalar fields for a useful one-row report.
            row = {k: v for k, v in data.items() if isinstance(v, (str, int, float, bool)) or v is None}
            return source.stem, ([row] if row else []), data
        if isinstance(data, list) and all(isinstance(x, dict) for x in data):
            return source.stem, data, {"source": source.name}
        return source.stem, [{"value": str(data)}], {"source": source.name}
    if source.suffix.lower() == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        return source.stem, rows, {"source": source.name, "row_count": len(rows)}
    raise ValueError("Choose a JSON or CSV experiment/report data file.")


def _numeric_series(rows: list[dict[str, Any]]) -> tuple[str | None, list[float]]:
    if not rows:
        return None, []
    excluded = {"seed", "generation", "episode", "step", "index"}
    for key in rows[0].keys():
        if key.lower() in excluded:
            continue
        values = []
        for row in rows:
            try:
                value = float(row.get(key, ""))
                if math.isfinite(value):
                    values.append(value)
                else:
                    values.append(float("nan"))
            except (TypeError, ValueError):
                values.append(float("nan"))
        finite = [x for x in values if math.isfinite(x)]
        if len(finite) >= 2:
            return key, values
    return None, []


def create_report_plot(source: Path, output_png: Path) -> Path | None:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return None
    title, rows, _ = load_report_data(source)
    key, values = _numeric_series(rows)
    if not key:
        return None
    xs = list(range(1, len(values) + 1))
    ys = [v if math.isfinite(v) else float("nan") for v in values]
    output_png.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(9, 4.5), constrained_layout=True)
    ax.plot(xs, ys, marker="o", linewidth=1.5)
    ax.set_title(f"{title}: {key}")
    ax.set_xlabel("Record")
    ax.set_ylabel(key.replace("_", " "))
    ax.grid(True, alpha=0.25)
    fig.savefig(output_png, dpi=180)
    plt.close(fig)
    return output_png


def export_report(source: Path, output: Path, fmt: str) -> Path:
    """Export JSON/CSV data into PDF, DOCX, or PPTX with an optional numeric plot."""
    source, output = Path(source), Path(output)
    title, rows, metadata = load_report_data(source)
    fmt = fmt.lower().lstrip(".")
    if fmt not in {"pdf", "docx", "pptx"}:
        raise ValueError("Output format must be PDF, DOCX, or PPTX.")
    output.parent.mkdir(parents=True, exist_ok=True)
    plot_path = output.with_name(output.stem + "_plot.png")
    plot = create_report_plot(source, plot_path)
    summary_lines = [
        f"Source file: {source.name}",
        f"Records: {len(rows)}",
        f"Generated (UTC): {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    for key, value in metadata.items():
        if isinstance(value, (str, int, float, bool)) and key not in {"history", "episodes", "results"}:
            summary_lines.append(f"{key}: {value}")
    preview_rows = rows[:30]
    if fmt == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import (
                Image,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError as exc:
            raise RuntimeError("PDF export requires reportlab. Install requirements-format-support.txt.") from exc
        doc = SimpleDocTemplate(str(output), pagesize=landscape(A4), title=title)
        styles = getSampleStyleSheet()
        story = [Paragraph(title.replace("&", "&amp;"), styles["Title"]), Spacer(1, 10)]
        story.extend(Paragraph(line.replace("&", "&amp;"), styles["BodyText"]) for line in summary_lines)
        story.append(Spacer(1, 12))
        if preview_rows:
            keys = list(dict.fromkeys(k for row in preview_rows for k in row.keys()))[:8]
            table_data = [keys]
            for row in preview_rows:
                table_data.append([str(row.get(k, ""))[:80] for k in keys])
            table = Table(table_data, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#263746")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
        if plot and plot.exists():
            story.extend([Spacer(1, 14), Image(str(plot), width=650, height=325)])
        doc.build(story)
    elif fmt == "docx":
        try:
            from docx import Document
            from docx.shared import Inches
        except ImportError as exc:
            raise RuntimeError("DOCX export requires python-docx. Install requirements-format-support.txt.") from exc
        doc = Document()
        doc.add_heading(title, 0)
        for line in summary_lines:
            doc.add_paragraph(line)
        if preview_rows:
            keys = list(dict.fromkeys(k for row in preview_rows for k in row.keys()))[:8]
            table = doc.add_table(rows=1, cols=max(1, len(keys)))
            table.style = "Light Shading Accent 1"
            for i, key in enumerate(keys):
                table.rows[0].cells[i].text = str(key)
            for row in preview_rows:
                cells = table.add_row().cells
                for i, key in enumerate(keys):
                    cells[i].text = str(row.get(key, ""))[:300]
        if plot and plot.exists():
            doc.add_heading("Numeric trend", level=1)
            doc.add_picture(str(plot), width=Inches(6.3))
        doc.save(output)
    else:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("PPTX export requires python-pptx. Install requirements-format-support.txt.") from exc
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "\n".join(summary_lines[:8])
        if plot and plot.exists():
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Numeric trend"
            slide.shapes.add_picture(str(plot), Inches(0.6), Inches(1.3), width=Inches(8.5))
        if preview_rows:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Data preview (first 15 records)"
            keys = list(dict.fromkeys(k for row in preview_rows[:15] for k in row.keys()))[:6]
            table = slide.shapes.add_table(
                rows=min(16, len(preview_rows) + 1),
                cols=max(1, len(keys)),
                left=Inches(0.3), top=Inches(1.3),
                width=Inches(9.0), height=Inches(5.2),
            ).table
            for j, key in enumerate(keys):
                table.cell(0, j).text = str(key)
            for i, row in enumerate(preview_rows[:15], start=1):
                for j, key in enumerate(keys):
                    table.cell(i, j).text = str(row.get(key, ""))[:60]
        prs.save(output)
    return output
